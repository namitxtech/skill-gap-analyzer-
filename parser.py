"""Resume text extraction (PDF / DOCX / TXT) and skill detection."""
import io
import re

from skills_data import CATALOG


def read_resume(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _pattern(alias: str) -> re.Pattern:
    # Letters, digits, + and # count as part of a word, so "java" won't match "javascript".
    return re.compile(r"(?<![A-Za-z0-9+#])" + re.escape(alias) + r"(?![A-Za-z0-9+#])", re.IGNORECASE)


_PATTERNS = {skill: [_pattern(a) for a in {skill.lower(), *meta[1]}] for skill, meta in CATALOG.items()}


def extract_skills(text: str) -> set[str]:
    """Return catalog skills mentioned anywhere in the text."""
    found = set()
    for skill, patterns in _PATTERNS.items():
        if any(p.search(text) for p in patterns):
            found.add(skill)
    return found
